#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Planilha de gestão de tráfego + funil (xlsx) com fórmulas e semáforo."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import os

OUT="/home/user/Test/06-gestao-kpis/Planilha-Gestao-Trafego-e-Funil.xlsx"
DARK="14110E"; GOLD="D8A84B"; GOLD2="F0CB6E"; CREAM="F4EEE2"; PANEL="2b2620"
GREEN="4CB286"; RED="C5524A"; YEL="E0A93B"; LIGHT="FAF6EC"

wb=openpyxl.Workbook()
thin=Side(style="thin", color="DDD3BE")
border=Border(left=thin,right=thin,top=thin,bottom=thin)
def hdr(c): c.fill=PatternFill("solid",fgColor=DARK); c.font=Font(color=CREAM,bold=True,size=10); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=border
def title(ws,cell,txt,size=15): ws[cell]=txt; ws[cell].font=Font(color=DARK,bold=True,size=size)
def money(c): c.number_format='R$ #,##0'
def pct(c): c.number_format='0.0%'

# =================================================================
# METAS
# =================================================================
ws=wb.active; ws.title="Metas"
title(ws,"A1","🎯 METAS / KPIs DE REFERÊNCIA")
ws["A2"]="Edite os alvos conforme sua praça e ticket. O Painel usa estes valores no semáforo."
ws["A2"].font=Font(italic=True,color="6A6151",size=9)
metas=[("Indicador","Alvo","Formato"),
 ("CPL alvo (custo por lead)",12,"R$"),
 ("Custo por conversa qualificada",60,"R$"),
 ("Custo por visita agendada",300,"R$"),
 ("Taxa Lead → Visita",0.20,"%"),
 ("Taxa Visita → Proposta",0.40,"%"),
 ("Taxa Proposta → Venda",0.45,"%"),
 ("Comissão média por venda",12000,"R$"),
 ("ROAS alvo (retorno s/ investimento)",5,"x")]
for r,(a,b,c) in enumerate(metas,4):
    ws.cell(r,1,a); ws.cell(r,2,b); ws.cell(r,3,c)
    if r==4:
        for col in (1,2,3): hdr(ws.cell(r,col))
    else:
        ws.cell(r,1).border=border; ws.cell(r,2).border=border; ws.cell(r,3).border=border
        if c=="R$": money(ws.cell(r,2))
        elif c=="%": pct(ws.cell(r,2))
        ws.cell(r,1).fill=PatternFill("solid",fgColor=LIGHT)
ws.column_dimensions["A"].width=38; ws.column_dimensions["B"].width=14; ws.column_dimensions["C"].width=10
# named cells for reference
M={"cpl":"Metas!$B$5","cconv":"Metas!$B$6","cvis":"Metas!$B$7","lv":"Metas!$B$8",
   "vp":"Metas!$B$9","pv":"Metas!$B$10","com":"Metas!$B$11","roas":"Metas!$B$12"}

# =================================================================
# LANÇAMENTOS (registro por campanha/dia)
# =================================================================
ws=wb.create_sheet("Lançamentos")
title(ws,"A1","📊 LANÇAMENTOS DE CAMPANHA (registre dia a dia)")
cols=["Data","Plataforma","Campanha","Etapa","Investido","Impressões","Cliques","Leads",
      "Conversas qualif.","Visitas agend.","Propostas","Vendas","CTR","CPL","Custo/Conversa"]
hr=3
for i,c in enumerate(cols,1):
    cell=ws.cell(hr,i,c); hdr(cell)
# dropdowns
dv_plat=DataValidation(type="list",formula1='"Meta,Google,YouTube,TikTok"',allow_blank=True)
dv_etapa=DataValidation(type="list",formula1='"TOFU,Retargeting,BOFU,Escala"',allow_blank=True)
ws.add_data_validation(dv_plat); ws.add_data_validation(dv_etapa)
# example data
ex=[
 ("2026-06-01","Meta","AQV|Meta|Msg|TOFU|Maceió","TOFU",60,18000,420,38,9,3,1,0),
 ("2026-06-01","Meta","AQV|Meta|Leads|TOFU|MCMV","TOFU",40,15000,300,52,7,2,1,0),
 ("2026-06-01","Google","AQV|Google|Search|BOFU","BOFU",50,2200,140,12,6,3,2,1),
 ("2026-06-02","Meta","AQV|Meta|Msg|TOFU|Maceió","TOFU",60,19500,455,41,11,4,1,1),
 ("2026-06-02","Meta","AQV|Meta|Msg|RMKT|Prova","Retargeting",25,8000,260,18,9,4,2,1),
 ("2026-06-02","Google","AQV|Google|Search|BOFU","BOFU",50,2400,150,14,7,4,2,1),
 ("2026-06-03","Meta","AQV|Meta|Msg|TOFU|Investidor","TOFU",40,12000,300,22,8,2,1,0),
 ("2026-06-03","TikTok","AQV|TikTok|Leads|TOFU|MCMV","TOFU",40,40000,900,80,6,1,0,0),
 ("2026-06-03","Meta","AQV|Meta|Msg|BOFU|Fechamento","BOFU",20,3000,120,10,7,5,3,2),
 ("2026-06-04","Meta","AQV|Meta|Msg|TOFU|Maceió","TOFU",70,21000,500,45,12,5,2,1),
 ("2026-06-04","Google","AQV|Google|PMax|Escala","Escala",50,9000,210,20,8,3,1,1),
 ("2026-06-04","YouTube","AQV|Google|DemandGen|RMKT","Retargeting",30,30000,180,9,4,2,1,0),
]
r=hr+1
for row in ex:
    for i,v in enumerate(row,1):
        c=ws.cell(r,i,v); c.border=border
        if i==5: money(c)
    # formulas CTR, CPL, Custo/Conversa
    ws.cell(r,13,f"=IFERROR(G{r}/F{r},0)"); pct(ws.cell(r,13))
    ws.cell(r,14,f"=IFERROR(E{r}/H{r},0)"); money(ws.cell(r,14))
    ws.cell(r,15,f"=IFERROR(E{r}/I{r},0)"); money(ws.cell(r,15))
    for i in (13,14,15): ws.cell(r,i).border=border
    dv_plat.add(ws.cell(r,2)); dv_etapa.add(ws.cell(r,4))
    r+=1
LAST=r-1
# add blank rows with formulas for future entries
for r in range(r, r+40):
    for i in range(1,13): ws.cell(r,i).border=border
    ws.cell(r,13,f"=IFERROR(G{r}/F{r},0)"); pct(ws.cell(r,13))
    ws.cell(r,14,f"=IFERROR(E{r}/H{r},0)"); money(ws.cell(r,14))
    ws.cell(r,15,f"=IFERROR(E{r}/I{r},0)"); money(ws.cell(r,15))
    money(ws.cell(r,5))
    dv_plat.add(ws.cell(r,2)); dv_etapa.add(ws.cell(r,4))
DATA_END=r
widths=[12,11,30,13,12,12,10,9,13,13,11,9,9,12,14]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes="A4"
# semáforo no CPL (col N) e Custo/Conversa (col O)
ws.conditional_formatting.add(f"N4:N{DATA_END}",
   CellIsRule(operator="lessThanOrEqual",formula=[M["cpl"]],fill=PatternFill("solid",fgColor="D6F0E4")))
ws.conditional_formatting.add(f"N4:N{DATA_END}",
   CellIsRule(operator="greaterThan",formula=[f"{M['cpl']}*1.5"],fill=PatternFill("solid",fgColor="F6D9D6")))
ws.conditional_formatting.add(f"O4:O{DATA_END}",
   CellIsRule(operator="lessThanOrEqual",formula=[M["cconv"]],fill=PatternFill("solid",fgColor="D6F0E4")))
ws.conditional_formatting.add(f"O4:O{DATA_END}",
   CellIsRule(operator="greaterThan",formula=[f"{M['cconv']}*1.5"],fill=PatternFill("solid",fgColor="F6D9D6")))

# =================================================================
# PAINEL (dashboard)
# =================================================================
ws=wb.create_sheet("Painel"); wb.move_sheet("Painel",-(len(wb.sheetnames)-1))  # move to first
title(ws,"A1","📈 PAINEL DE PERFORMANCE")
ws["A2"]="Atualiza automaticamente conforme você preenche a aba Lançamentos."
ws["A2"].font=Font(italic=True,color="6A6151",size=9)
L="Lançamentos"
rng=lambda col: f"{L}!{col}4:{col}{DATA_END}"
kpis=[
 ("Investimento total",f"=SUM({rng('E')})","R$"),
 ("Leads",f"=SUM({rng('H')})","int"),
 ("Conversas qualificadas",f"=SUM({rng('I')})","int"),
 ("Visitas agendadas",f"=SUM({rng('J')})","int"),
 ("Propostas",f"=SUM({rng('K')})","int"),
 ("Vendas",f"=SUM({rng('L')})","int"),
 ("CPL médio",f"=IFERROR(SUM({rng('E')})/SUM({rng('H')}),0)","R$"),
 ("Custo por conversa",f"=IFERROR(SUM({rng('E')})/SUM({rng('I')}),0)","R$"),
 ("Custo por visita",f"=IFERROR(SUM({rng('E')})/SUM({rng('J')}),0)","R$"),
 ("Custo por venda (CAC)",f"=IFERROR(SUM({rng('E')})/SUM({rng('L')}),0)","R$"),
 ("Taxa Lead → Visita",f"=IFERROR(SUM({rng('J')})/SUM({rng('H')}),0)","%"),
 ("Taxa Visita → Proposta",f"=IFERROR(SUM({rng('K')})/SUM({rng('J')}),0)","%"),
 ("Taxa Proposta → Venda",f"=IFERROR(SUM({rng('L')})/SUM({rng('K')}),0)","%"),
 ("Receita (comissão)",f"=SUM({rng('L')})*{M['com']}","R$"),
 ("ROAS (retorno/investido)",f"=IFERROR((SUM({rng('L')})*{M['com']})/SUM({rng('E')}),0)","x"),
]
r=4
hdr(ws.cell(r,1,"Indicador")); hdr(ws.cell(r,2,"Valor")); hdr(ws.cell(r,3,"Meta")); hdr(ws.cell(r,4,"Status"))
metas_map={"CPL médio":M["cpl"],"Custo por conversa":M["cconv"],"Custo por visita":M["cvis"],
 "Taxa Lead → Visita":M["lv"],"Taxa Visita → Proposta":M["vp"],"Taxa Proposta → Venda":M["pv"],
 "ROAS (retorno/investido)":M["roas"]}
lower_is_better={"CPL médio","Custo por conversa","Custo por visita"}
r=5
for (nome,form,fmt) in kpis:
    ws.cell(r,1,nome).font=Font(bold=True,size=10)
    c=ws.cell(r,2,form)
    if fmt=="R$": money(c)
    elif fmt=="%": pct(c)
    elif fmt=="x": c.number_format='0.0"x"'
    else: c.number_format='#,##0'
    ws.cell(r,1).border=border; c.border=border
    ws.cell(r,1).fill=PatternFill("solid",fgColor=LIGHT)
    if nome in metas_map:
        mc=ws.cell(r,3,f"={metas_map[nome]}")
        if fmt=="R$": money(mc)
        elif fmt=="%": pct(mc)
        elif fmt=="x": mc.number_format='0.0"x"'
        mc.border=border
        if nome in lower_is_better:
            ws.cell(r,4,f'=IF(B{r}=0,"—",IF(B{r}<=C{r},"✅ no alvo",IF(B{r}<=C{r}*1.5,"⚠️ atenção","🔴 caro")))')
        else:
            ws.cell(r,4,f'=IF(B{r}=0,"—",IF(B{r}>=C{r},"✅ no alvo",IF(B{r}>=C{r}*0.7,"⚠️ atenção","🔴 baixo")))')
        ws.cell(r,4).border=border; ws.cell(r,4).alignment=Alignment(horizontal="center")
    r+=1
for i,w in enumerate([30,16,14,16],1): ws.column_dimensions[get_column_letter(i)].width=w
# color the status column
ws.conditional_formatting.add(f"D5:D{r}",
   FormulaRule(formula=['ISNUMBER(SEARCH("✅",D5))'],fill=PatternFill("solid",fgColor="D6F0E4")))
ws.conditional_formatting.add(f"D5:D{r}",
   FormulaRule(formula=['ISNUMBER(SEARCH("⚠️",D5))'],fill=PatternFill("solid",fgColor="FCEFCB")))
ws.conditional_formatting.add(f"D5:D{r}",
   FormulaRule(formula=['ISNUMBER(SEARCH("🔴",D5))'],fill=PatternFill("solid",fgColor="F6D9D6")))

# =================================================================
# FUNIL (pipeline de leads / atendimentos)
# =================================================================
ws=wb.create_sheet("Funil de Vendas")
title(ws,"A1","🧭 FUNIL DE VENDAS (atendimentos atuais)")
fcols=["Lead","Origem","Entrada","Estágio","Imóvel","Valor imóvel","Comissão potencial","Última interação","Dias parado","Próxima ação"]
fr=3
for i,c in enumerate(fcols,1): hdr(ws.cell(fr,i,c))
dv_orig=DataValidation(type="list",formula1='"Meta,Google,YouTube,TikTok,Indicação,Orgânico"',allow_blank=True)
dv_est=DataValidation(type="list",formula1='"Novo,Conversa iniciada,Qualificado,Visita agendada,Proposta,Venda,Perdido"',allow_blank=True)
ws.add_data_validation(dv_orig); ws.add_data_validation(dv_est)
fex=[
 ("Marina S.","Meta","2026-06-01","Visita agendada","Apê 2q [Bairro]",389000,"=F4*0.05","2026-06-04","=TODAY()-H4","Confirmar visita sáb"),
 ("João P.","Google","2026-06-02","Proposta","Cobertura frente-mar",1200000,"=F5*0.05","2026-06-05","=TODAY()-H5","Negociar contraproposta"),
 ("Ana e Léo","Meta","2026-06-02","Qualificado","Apê 3q upgrade",560000,"=F6*0.05","2026-06-03","=TODAY()-H6","Enviar simulação"),
 ("Carlos M.","TikTok","2026-06-03","Conversa iniciada","MCMV 2q",250000,"=F7*0.05","2026-06-03","=TODAY()-H7","Qualificar renda/FGTS"),
 ("Investidor RJ","Meta","2026-06-03","Qualificado","Studio frente-mar",430000,"=F8*0.05","2026-06-04","=TODAY()-H8","Enviar estudo ROI"),
]
r=fr+1
for row in fex:
    for i,v in enumerate(row,1):
        c=ws.cell(r,i,v); c.border=border
        if i in (6,7): money(c)
    dv_orig.add(ws.cell(r,2)); dv_est.add(ws.cell(r,4)); r+=1
for r in range(r,r+30):
    for i in range(1,11): ws.cell(r,i).border=border
    ws.cell(r,7,f"=IFERROR(F{r}*0.05,0)"); money(ws.cell(r,7)); money(ws.cell(r,6))
    ws.cell(r,9,f'=IF(H{r}="","",TODAY()-H{r})')
    dv_orig.add(ws.cell(r,2)); dv_est.add(ws.cell(r,4))
FEND=r
for i,w in enumerate([16,11,12,18,20,15,16,15,11,26],1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes="A4"
# highlight leads parados > 3 dias (vermelho) e estágios
ws.conditional_formatting.add(f"I4:I{FEND}",
   CellIsRule(operator="greaterThan",formula=["3"],fill=PatternFill("solid",fgColor="F6D9D6"),font=Font(color="9A2A22",bold=True)))
ws.conditional_formatting.add(f"D4:D{FEND}",
   FormulaRule(formula=['$D4="Venda"'],fill=PatternFill("solid",fgColor="D6F0E4")))
ws.conditional_formatting.add(f"D4:D{FEND}",
   FormulaRule(formula=['$D4="Perdido"'],fill=PatternFill("solid",fgColor="EDE7DA")))

# resumo do funil (contagem por estágio)
ws.cell(2,8,"Resumo do funil →").font=Font(bold=True,color=GOLD.replace('D8A84B','9a6f17'))
stages=["Conversa iniciada","Qualificado","Visita agendada","Proposta","Venda","Perdido"]
ws.cell(fr,12,"Estágio"); ws.cell(fr,13,"Qtd"); hdr(ws.cell(fr,12)); hdr(ws.cell(fr,13))
for i,s in enumerate(stages,1):
    ws.cell(fr+i,12,s).border=border
    ws.cell(fr+i,13,f'=COUNTIF($D$4:$D${FEND},"{s}")').border=border
ws.column_dimensions["L"].width=18; ws.column_dimensions["M"].width=8

wb.save(OUT)
print("Planilha salva:",OUT)
print("Abas:",wb.sheetnames)
